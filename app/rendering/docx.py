"""Render a document model as a Word file.

Word is what a person can actually edit before sending, which is why it is
offered alongside the PDF. Built-in heading styles, a document title and a
language are set so the file is navigable in a screen reader, and a table's
header row is marked to repeat across pages.
"""

import io

import docx
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import _Row
from docx.text.paragraph import Paragraph

from app.domain.document_model import Block, BlockType, DocumentModel, Run

#: A4 in millimetres.
PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
MARGIN_MM = 20

#: python-docx maps these to the built-in styles a screen reader navigates by.
_HEADING_STYLES = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}

#: Fonts that cover Devanagari, so a Hindi or Marathi letter renders correctly.
_LATIN_FONT = "Noto Serif"
_COMPLEX_FONT = "Noto Serif Devanagari"


def render(document: DocumentModel) -> bytes:
    """Render the document model as a ``.docx`` file.

    Args:
        document: The document model.

    Returns:
        The file's bytes.
    """
    word = docx.Document()
    _set_page(word)
    _set_metadata(word, document)
    _set_language(word, document.language)

    for block in document.blocks:
        _add_block(word, block)

    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _set_page(word: WordDocument) -> None:
    """Set A4 with even margins."""
    for section in word.sections:
        section.page_width = Mm(PAGE_WIDTH_MM)
        section.page_height = Mm(PAGE_HEIGHT_MM)
        section.left_margin = section.right_margin = Mm(MARGIN_MM)
        section.top_margin = section.bottom_margin = Mm(MARGIN_MM)


def _set_metadata(word: WordDocument, document: DocumentModel) -> None:
    """Set the title, which assistive technology announces when the file opens."""
    word.core_properties.title = document.title
    word.core_properties.language = document.language
    style = word.styles["Normal"]
    style.font.name = _LATIN_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:cs"), _COMPLEX_FONT)


def _set_language(word: WordDocument, language: str) -> None:
    """Tag the default style with the document's language."""
    rpr = word.styles["Normal"].element.get_or_add_rPr()
    lang = rpr.makeelement(qn("w:lang"), {})
    lang.set(qn("w:val"), language)
    lang.set(qn("w:bidi"), language)
    rpr.append(lang)


def _add_block(word: WordDocument, block: Block) -> None:
    """Append one block to the Word document."""
    if block.type is BlockType.HEADING:
        paragraph = word.add_paragraph(style=_HEADING_STYLES.get(block.level, "Heading 2"))
        _add_runs(paragraph, block.runs)
    elif block.type is BlockType.LIST:
        style = "List Number" if block.ordered else "List Bullet"
        for item in block.items:
            word.add_paragraph(item, style=style)
    elif block.type is BlockType.QUOTE:
        paragraph = word.add_paragraph(style="Intense Quote")
        _add_runs(paragraph, block.runs)
    elif block.type is BlockType.TABLE:
        _add_table(word, block)
    elif block.type is BlockType.RULE:
        paragraph = word.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("* * *")
    else:
        _add_runs(word.add_paragraph(), block.runs)


def _add_runs(paragraph: Paragraph, runs: list[Run]) -> None:
    """Append formatted runs to a paragraph."""
    for run in runs:
        added = paragraph.add_run(run.text)
        added.bold = run.bold
        added.italic = run.italic


def _add_table(word: WordDocument, block: Block) -> None:
    """Append a table, marking the header row to repeat on every page."""
    if not block.rows:
        return
    table = word.add_table(rows=0, cols=len(block.rows[0]))
    table.style = "Table Grid"

    for index, row in enumerate(block.rows):
        cells = table.add_row().cells
        for cell, source in zip(cells, row, strict=False):
            cell.text = ""
            _add_runs(cell.paragraphs[0], source.runs)
        if index == 0 and any(source.header for source in row):
            _repeat_header(table.rows[0])


def _repeat_header(row: _Row) -> None:
    """Mark a table row as a header row that repeats across page breaks."""
    properties = row._tr.get_or_add_trPr()
    header = properties.makeelement(qn("w:tblHeader"), {})
    header.set(qn("w:val"), "true")
    properties.append(header)

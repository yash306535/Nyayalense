"""Letter drafting, the slot lock, the fact audit and the exported files."""

import io

import pdfplumber
import pytest
from docx import Document as WordDocument
from httpx import AsyncClient

FACTS = {
    "sender_name": "Kavita Rane",
    "sender_address": "22 Model Colony, Pune",
    "sender_email": "kavita@example.com",
    "recipient_name": "Anil Deshpande",
    "recipient_address": "14 Shivaji Nagar, Pune",
    "premises": "Flat 7B, Sunrise Residency",
    "agreement_date": "2026-04-01",
    "deposit_amount": "60,000",
    "move_out_date": "2027-03-31",
    "bank_details": "UPI kavita@bank",
    "additional_points": "Please confirm receipt.",
    "reply_days": "15",
}


async def test_templates_declare_their_own_fields(client: AsyncClient) -> None:
    templates = (await client.get("/api/v1/drafts/templates")).json()
    assert {template["id"] for template in templates} == {
        "deposit_refund_request",
        "clause_change_request",
        "resignation_letter",
        "grievance_letter",
    }
    for template in templates:
        assert template["fields"]
        for field in template["fields"]:
            assert field["labels"]["en"]


async def test_no_template_calls_itself_a_legal_notice(client: AsyncClient) -> None:
    """Titles must not imply the writer has a lawyer behind them."""
    templates = (await client.get("/api/v1/drafts/templates")).json()
    for template in templates:
        assert "legal notice" not in template["title"].casefold()


async def test_templates_can_be_narrowed_to_a_document_type(client: AsyncClient) -> None:
    response = await client.get("/api/v1/drafts/templates", params={"doc_type": "employment_offer"})
    assert "resignation_letter" in {template["id"] for template in response.json()}


@pytest.mark.parametrize(
    ("template_id", "facts"),
    [
        ("deposit_refund_request", FACTS),
        (
            "resignation_letter",
            {
                "sender_name": "R Sharma",
                "designation": "Engineer",
                "recipient_name": "HR",
                "company_name": "Meridian",
                "notice_period_days": "90",
                "resignation_date": "2026-09-01",
                "last_working_day": "2026-11-30",
                "settlement_days": "45",
            },
        ),
        (
            "grievance_letter",
            {
                "sender_name": "R Sharma",
                "sender_address": "Pune",
                "company_name": "Insurer",
                "account_reference": "POL-1",
                "incident_date": "2026-08-01",
                "what_happened": "The claim was refused.",
                "what_you_want": "Reconsider it.",
                "reply_days": "15",
            },
        ),
        (
            "clause_change_request",
            {
                "sender_name": "R Sharma",
                "recipient_name": "Owner",
                "document_name": "Agreement",
                "clause_label": "5.2",
                "clause_quote": "The deposit shall stand forfeited.",
                "requested_change": "Remove the forfeiture.",
                "reply_days": "7",
            },
        ),
    ],
)
async def test_every_template_renders(
    client: AsyncClient, template_id: str, facts: dict[str, str]
) -> None:
    response = await client.post(
        "/api/v1/drafts/preview", json={"template_id": template_id, "facts": facts}
    )
    assert response.status_code == 200, response.text
    assert response.json()["errors"] == {}
    assert response.json()["document"]["blocks"]


async def test_a_missing_required_fact_is_reported_per_field(client: AsyncClient) -> None:
    response = await client.post(
        "/api/v1/drafts/preview",
        json={"template_id": "deposit_refund_request", "facts": {"sender_name": "R"}},
    )
    assert response.status_code == 200
    assert "deposit_amount" in response.json()["errors"]


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("agreement_date", "not-a-date"),
        ("agreement_date", "2026-02-30"),
        ("sender_email", "not an email"),
        ("deposit_amount", "sixty thousand"),
        ("reply_days", "99"),
    ],
)
async def test_badly_formatted_values_are_rejected(
    client: AsyncClient, field: str, value: str
) -> None:
    response = await client.post(
        "/api/v1/drafts/preview",
        json={"template_id": "deposit_refund_request", "facts": {**FACTS, field: value}},
    )
    assert field in response.json()["errors"]


async def test_markdown_inside_a_fact_prints_literally(client: AsyncClient) -> None:
    """A fact must never become document structure."""
    response = await client.post(
        "/api/v1/drafts/preview",
        json={
            "template_id": "deposit_refund_request",
            "facts": {**FACTS, "additional_points": "# Not a heading and **not bold**"},
        },
    )
    document = response.json()["document"]
    headings = [block for block in document["blocks"] if block["type"] == "heading"]
    assert all("Not a heading" not in block["runs"][0]["text"] for block in headings)
    assert "# Not a heading and **not bold**" in _text(document)


async def test_an_unknown_template_is_not_found(client: AsyncClient) -> None:
    response = await client.post(
        "/api/v1/drafts/preview", json={"template_id": "nope", "facts": {}}
    )
    assert response.status_code == 404


async def test_prefill_only_offers_what_a_verified_result_contains(
    client: AsyncClient, rental_document
) -> None:
    body = {
        "document": rental_document.model_dump(mode="json"),
        "audience": {"role": "tenant", "language": "en", "reading_level": "simple"},
    }
    overview = (await client.post("/api/v1/analysis/overview", json=body)).json()

    response = await client.post(
        "/api/v1/drafts/prefill",
        json={
            "template_id": "deposit_refund_request",
            "document": body["document"],
            "overview": overview,
        },
    )
    assert response.status_code == 200
    prefilled = response.json()
    for name in prefilled["sources"]:
        assert name in prefilled["values"], "a source without a value would be meaningless"
        assert prefilled["sources"][name]["label"]


# ---------------------------------------------------------------- exports


def _text(document: dict[str, object]) -> str:
    parts = []
    for block in document["blocks"]:
        parts.extend(run["text"] for run in block.get("runs", []))
        parts.extend(block.get("items", []))
    return " ".join(parts)


async def export(client: AsyncClient, export_format: str, *, confirmed: bool = True):
    return await client.post(
        "/api/v1/exports",
        json={
            "kind": "draft",
            "format": export_format,
            "template_id": "deposit_refund_request",
            "facts": FACTS,
            "confirmed": confirmed,
        },
    )


async def test_a_draft_cannot_be_downloaded_before_the_facts_are_confirmed(
    client: AsyncClient,
) -> None:
    response = await export(client, "pdf", confirmed=False)
    assert response.status_code == 422
    assert response.json()["type"].endswith("/facts-not-confirmed")


async def test_the_word_file_is_a4_with_real_headings_and_every_fact(
    client: AsyncClient,
) -> None:
    response = await export(client, "docx")
    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith("attachment;")

    word = WordDocument(io.BytesIO(response.content))
    assert round(word.sections[0].page_width.mm) == 210
    assert round(word.sections[0].page_height.mm) == 297
    assert word.core_properties.title
    assert any(p.style.name.startswith("Heading") for p in word.paragraphs)

    text = "\n".join(p.text for p in word.paragraphs)
    for value in ("Kavita Rane", "Anil Deshpande", "60,000", "2027-03-31"):
        assert value in text


async def test_a_word_table_header_row_repeats(client: AsyncClient, rental_document) -> None:
    body = {
        "document": rental_document.model_dump(mode="json"),
        "audience": {"role": "tenant", "language": "en", "reading_level": "simple"},
    }
    overview = (await client.post("/api/v1/analysis/overview", json=body)).json()
    response = await client.post(
        "/api/v1/exports",
        json={"kind": "brief", "format": "docx", "brief": {"title": "T", "overview": overview}},
    )
    word = WordDocument(io.BytesIO(response.content))
    assert word.tables
    assert "tblHeader" in word.tables[0].rows[0]._tr.xml


async def test_the_pdf_is_a4_with_metadata_and_every_fact(client: AsyncClient) -> None:
    response = await export(client, "pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")

    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        page = pdf.pages[0]
        assert 592 < page.width < 598, "A4 width in points"
        assert 839 < page.height < 845, "A4 height in points"
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        assert pdf.metadata.get("Title")

    for value in ("Kavita Rane", "60,000"):
        assert value in text


async def test_the_filename_is_dated_and_safe(client: AsyncClient) -> None:
    disposition = (await export(client, "docx")).headers["content-disposition"]
    assert ".." not in disposition
    assert "/" not in disposition.split("filename=")[1]
    assert ".docx" in disposition


async def test_exporting_an_unknown_kind_is_refused(client: AsyncClient) -> None:
    response = await client.post("/api/v1/exports", json={"kind": "nope", "format": "pdf"})
    assert response.status_code == 422


async def test_a_brief_export_needs_a_brief(client: AsyncClient) -> None:
    response = await client.post("/api/v1/exports", json={"kind": "brief", "format": "pdf"})
    assert response.status_code == 422
